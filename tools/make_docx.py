"""生成 2026 MathorCup 大数据竞赛论文骨架（.docx）。

对齐官方 Word 模板：
  A4 / 上下 2.0cm 左右 2.5cm / 正文小四(12pt) / 封面(编号·赛道·题目)
  / 摘要页 / 目录页(自动域) / 正文 / 页脚「第 X 页 共 Y 页」

用法：
  python make_docx.py 输出路径.docx "队伍编号" "A" "论文题目"
"""
from __future__ import annotations

import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CN_FONT = "宋体"
CN_HEI = "黑体"
EN_FONT = "Times New Roman"


def set_run(run, size=12, bold=False, cn=CN_FONT, en=EN_FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = en
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn)
    run.font.color.rgb = RGBColor(0, 0, 0)


def para(doc, text="", size=12, bold=False, align=None, cn=CN_FONT,
         space_before=0, space_after=0, line=1.5):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if text:
        set_run(p.add_run(text), size=size, bold=bold, cn=cn)
    return p


def add_field(paragraph, instr: str):
    r = paragraph.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "1"
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
    for el in (f1, it, f2, t, f3):
        r._r.append(el)


def setup_page(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.75)


def footer_page_numbers(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("第 "), size=9)
    add_field(p, "PAGE")
    set_run(p.add_run(" 页  共 "), size=9)
    add_field(p, "NUMPAGES")
    set_run(p.add_run(" 页"), size=9)


def build(path: str, bianhao: str, saidao: str, timu: str):
    doc = Document()

    # 默认样式
    style = doc.styles["Normal"]
    style.font.name = EN_FONT
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

    sec = doc.sections[0]
    setup_page(sec)
    footer_page_numbers(sec)

    # ---------- 封面 ----------
    para(doc, "", space_before=72)
    para(doc, "2026 年 MathorCup 数学应用挑战赛", size=16, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "——大数据竞赛", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         space_after=48)

    para(doc, "队伍编号", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, bianhao, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)

    para(doc, "赛道", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, saidao, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)

    para(doc, "论文题目", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, timu, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_page_break()

    # ---------- 摘要 ----------
    para(doc, timu, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         space_after=12)
    para(doc, "摘　要", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         space_after=12)
    para(doc, "（摘要正文：600–800 字，逐问一段，写明方法名与关键数值结果。）",
         size=12, line=1.5)
    for i in range(1, 5):
        para(doc, "针对问题%s，……（方法、模型与量化结果）。" % "一二三四"[i - 1],
             size=12, line=1.5, space_before=6)
    para(doc, "", space_before=12)
    p = doc.add_paragraph()
    set_run(p.add_run("关键词："), size=12, bold=True)
    set_run(p.add_run("关键词1；关键词2；关键词3；关键词4"), size=12)

    doc.add_page_break()

    # ---------- 目录 ----------
    para(doc, "目　录", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         space_after=12)
    p = doc.add_paragraph()
    add_field(p, r'TOC \o "1-3" \h \z \u')
    para(doc, "（在 Word/WPS 中右键“更新域”即可生成目录；提交前务必更新一次。）",
         size=9)

    doc.add_page_break()

    # ---------- 正文 ----------
    sections = [
        ("一、问题重述", ["1.1 问题背景", "1.2 问题提出"]),
        ("二、问题分析", ["2.1 数据概况", "2.2 总体思路"]),
        ("三、模型假设", []),
        ("四、符号说明", []),
        ("五、问题一模型的建立与求解", ["5.1 数据预处理", "5.2 模型建立", "5.3 求解与结果分析"]),
        ("六、问题二模型的建立与求解", ["6.1 模型建立", "6.2 求解与结果分析"]),
        ("七、问题三模型的建立与求解", ["7.1 模型建立", "7.2 求解与结果分析"]),
        ("八、模型评价与推广", ["8.1 模型优点", "8.2 模型不足", "8.3 模型推广"]),
        ("参考文献", []),
        ("附录", ["A 支撑材料清单", "B 核心代码"]),
    ]
    for title, subs in sections:
        para(doc, title, size=14, bold=True, cn=CN_HEI,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=6)
        for s in subs:
            para(doc, s, size=12, bold=True, cn=CN_HEI, space_before=6)
            para(doc, "（正文）", size=12, line=1.5)

    doc.save(path)
    return path


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "MathorCup大数据赛论文模板.docx"
    bh = sys.argv[2] if len(sys.argv) > 2 else "MCB260XXXX"
    sd = sys.argv[3] if len(sys.argv) > 3 else "A"
    tm = sys.argv[4] if len(sys.argv) > 4 else "在此填写论文题目"
    print("已生成:", build(out, bh, sd, tm))
